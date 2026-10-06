import re
from pathlib import Path

from pipeline.pitch_schema import (
    PITCH_CRITERIA,
    create_pitch_result,
    validate_pitch_result,
)


def extract_text_from_file(file_path):
    """
    Extract text from a PDF or PPTX pitch deck.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf_text(path)

    if suffix == ".pptx":
        return _extract_pptx_text(path)

    raise ValueError(
        "Unsupported pitch deck format. "
        "Only PDF and PPTX are supported."
    )


def _extract_pdf_text(path):
    """
    Extract text from PDF using PyMuPDF.
    """

    try:
        import fitz
    except ImportError:
        raise ImportError(
            "PyMuPDF is required for PDF extraction. "
            "Install it using: pip install pymupdf"
        )

    document = fitz.open(path)
    pages = []

    try:
        for page in document:
            text = page.get_text("text")

            if text:
                pages.append(text)

    finally:
        document.close()

    return "\n".join(pages).strip()


def _extract_pptx_text(path):
    """
    Extract text from PPTX using python-pptx.
    """

    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError(
            "python-pptx is required for PPTX extraction. "
            "Install it using: pip install python-pptx"
        )

    presentation = Presentation(path)
    slides = []

    for slide in presentation.slides:
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_text.append(shape.text.strip())

        if slide_text:
            slides.append("\n".join(slide_text))

    return "\n\n".join(slides).strip()


def _normalize_text(text):
    """
    Normalize extracted text.
    """

    text = text.lower()
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def _find_evidence(text, keywords, max_items=3):
    """
    Find evidence sentences containing relevant keywords.
    """

    sentences = re.split(
        r"(?<=[.!?])\s+",
        text
    )

    evidence = []

    for sentence in sentences:
        normalized = sentence.lower()

        if any(
            keyword.lower() in normalized
            for keyword in keywords
        ):
            cleaned = sentence.strip()

            if cleaned and cleaned not in evidence:
                evidence.append(cleaned)

        if len(evidence) >= max_items:
            break

    return evidence


def _analyze_criterion(text, criterion):
    """
    Perform deterministic first-pass analysis
    for one pitch criterion.
    """

    keyword_map = {
        "Problem": [
            "problem",
            "pain point",
            "challenge",
            "gap",
            "difficulty",
        ],
        "Solution": [
            "solution",
            "product",
            "platform",
            "technology",
            "solve",
        ],
        "Market": [
            "tam",
            "sam",
            "som",
            "market size",
            "market",
            "industry",
        ],
        "Business_Model": [
            "business model",
            "revenue model",
            "subscription",
            "commission",
            "saas",
            "pricing",
            "revenue",
        ],
        "Traction": [
            "traction",
            "customers",
            "users",
            "revenue",
            "growth",
            "sales",
            "orders",
        ],
        "Competition": [
            "competitor",
            "competition",
            "competitive",
            "market leader",
            "alternative",
        ],
        "Product_Technology": [
            "technology",
            "platform",
            "software",
            "product",
            "ip",
            "patent",
            "technology readiness",
        ],
        "Team": [
            "founder",
            "co-founder",
            "team",
            "experience",
            "advisor",
            "leadership",
        ],
        "Financials": [
            "revenue",
            "ebitda",
            "profit",
            "margin",
            "burn",
            "cash",
            "financial",
        ],
        "Fundraising": [
            "fundraising",
            "funding",
            "raise",
            "investment",
            "ask",
            "valuation",
            "equity",
        ],
        "Scalability": [
            "scale",
            "scalable",
            "expansion",
            "growth",
            "geography",
            "international",
        ],
    }

    keywords = keyword_map.get(criterion, [])

    evidence = _find_evidence(
        text,
        keywords,
        max_items=3,
    )

    missing_information = []
    risks = []

    if evidence:
        score = min(
            100,
            40 + (len(evidence) * 15)
        )

        status = "Supported"
        confidence = "Medium"

    else:
        score = None
        status = "Missing"
        confidence = "Low"

    if not evidence:
        missing_information.append(
            f"No clear evidence found for {criterion}."
        )

    elif len(evidence) == 1:
        missing_information.append(
            f"Limited evidence available for {criterion}."
        )

    if criterion in {
        "Market",
        "Financials",
        "Traction",
        "Fundraising",
    } and not evidence:

        risks.append(
            f"{criterion} cannot be adequately assessed "
            "from the pitch deck."
        )

    return {
        "score": score,
        "evidence": evidence,
        "missing_information": missing_information,
        "risks": risks,
        "confidence": confidence,
        "evidence_status": status,
    }


def analyze_pitch_deck(file_path):
    """
    Analyze a pitch deck and return the
    standardized VentureIQ pitch-analysis result.
    """

    raw_text = extract_text_from_file(file_path)

    if not raw_text:
        raise ValueError(
            "No readable text was found in the pitch deck."
        )

    text = _normalize_text(raw_text)

    result = create_pitch_result()

    scores = []

    for criterion in PITCH_CRITERIA:

        criterion_result = _analyze_criterion(
            text,
            criterion,
        )

        result["criteria"][criterion] = criterion_result

        if criterion_result["score"] is not None:
            scores.append(
                criterion_result["score"]
            )

        if criterion_result["evidence"]:

            result["strengths"].append(
                f"{criterion} has supporting evidence "
                "in the pitch deck."
            )

        result["risks"].extend(
            criterion_result["risks"]
        )

        result["due_diligence_actions"].extend(
            criterion_result["missing_information"]
        )

    # Calculate overall pitch score.
    if scores:

        result["pitch_score"] = round(
            sum(scores) / len(scores),
            2,
        )

    # Calculate evidence confidence.
    confidence_scores = {
        "High": 100,
        "Medium": 60,
        "Low": 30,
    }

    confidence_values = []

    for criterion in PITCH_CRITERIA:

        confidence = result["criteria"][criterion][
            "confidence"
        ]

        if confidence:

            confidence_values.append(
                confidence_scores.get(
                    confidence,
                    0,
                )
            )

    if confidence_values:

        result["evidence_confidence"] = round(
            sum(confidence_values)
            / len(confidence_values),
            2,
        )

    # Remove duplicates while preserving order.
    result["strengths"] = list(
        dict.fromkeys(
            result["strengths"]
        )
    )

    result["risks"] = list(
        dict.fromkeys(
            result["risks"]
        )
    )

    result["due_diligence_actions"] = list(
        dict.fromkeys(
            result["due_diligence_actions"]
        )
    )

    # Validate final structure.
    validate_pitch_result(result)

    return result